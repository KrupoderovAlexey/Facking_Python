import openpyxl as op

# 1. Чтение файла Excel
wb = op.load_workbook(r"C:\Users\user\Desktop\Shit_EXEL.xlsx", read_only=True, data_only=True)
sh = wb.active
data = []
for row in sh.iter_rows(values_only=True):
    if all(c is None for c in row):
        break
    data.append(list(row))

print("Исходные данные из Excel:")
for i in range(len(data)):
    print(i, "   ", data[i])

# ====================================================================
# Обработка данных (начинаем с индекса 1, пропускаем шапку таблицы)
# ====================================================================
for i in range(1, len(data)):
    row = data[i]
    
    # ----------------------------------------------------------------
    # [Индекс 1] Даты в один единый формат ДД.ММ.ГГГГ
    # ----------------------------------------------------------------
    cl = str(row[1]).strip()
    if '-' in cl: parts = cl.split('-')
    elif '.' in cl: parts = cl.split('.')
    elif '/' in cl: parts = cl.split('/')
    elif ' ' in cl: parts = cl.split(' ')
    else: parts = [cl]
    
    if len(parts) == 3:
        if len(parts[0]) == 4: # Если год в начале (ГГГГ-ММ-ДД)
            row[1] = f"{parts[2]}.{parts[1]}.{parts[0]}"
        else: # Если день в начале (ДД.ММ.ГГГГ)
            row[1] = f"{parts[0]}.{parts[1]}.{parts[2]}"
    else:
        row[1] = cl.replace('/', '.')

    # ----------------------------------------------------------------
    # [Индекс 2] Товар в один формат (Каждое Слово С Заглавной Буквы)
    # ----------------------------------------------------------------
    row[2] = str(row[2]).strip().title()

    # ----------------------------------------------------------------
    # [Индекс 3] Количество в один формат int (Безопасно)
    # ----------------------------------------------------------------
    val_q = row[3]
    if isinstance(val_q, str):
        try:
            row[3] = int(val_q.strip())
        except ValueError:
            row[3] = 0 # Заменяем текст вроде 'не указана' на 0
    elif isinstance(val_q, (int, float)):
        row[3] = int(val_q)
    else:
        row[3] = 0

    # ----------------------------------------------------------------
    # [Индекс 4] Приведение цены к единому виду (Отрезаем буквы через isdigit)
    # ----------------------------------------------------------------
    val_p = row[4]
    if isinstance(val_p, str):
        num = val_p.replace(" ", "").strip()
        while num and not num.isdigit():
            num = num[:-1] # Отрезаем по одной букве ('р', 'у', 'б' и т.д.) с конца
        row[4] = int(num) if num.isdigit() else 0
    elif isinstance(val_p, (int, float)):
        row[4] = int(val_p)
    else:
        row[4] = 0

    # ----------------------------------------------------------------
    # [Индекс 5] Приведение к одному формату скидки (Переводим % в float)
    # ----------------------------------------------------------------
    val_d = row[5]
    if isinstance(val_d, str):
        p = val_d.strip()
        if 'без' in p or p == '':
            row[5] = 0.0
        elif '%' in p:
            p_clean = p.replace('%', '').strip()
            row[5] = round(float(p_clean) / 100, 2)
        else:
            try:
                row[5] = float(p)
            except ValueError:
                row[5] = 0.0
    elif isinstance(val_d, (int, float)):
        row[5] = float(val_d)
    else:
        row[5] = 0.0

    # ----------------------------------------------------------------
    # [Индекс 6] Приведение к одному формату статуса
    # ----------------------------------------------------------------
    row[6] = str(row[6]).strip().capitalize()

# ====================================================================
# Удаление повторов строк таблицы
# ====================================================================
new_data = [data[0]]
for row in data[1:]:
    if row not in new_data:
        new_data.append(row)
data = new_data

# Вывод результатов в консоль
print('\n\nИтоговые отформатированные данные:')
for i in range(len(data)):
    print(i, "   ", data[i])

# ====================================================================
# Запись результата в новый файл Excel
# ====================================================================
new_wb = op.Workbook()
ww = new_wb.active
for row in data:
    ww.append(row)

s = r"C:\Users\user\Desktop\MyAnswer.xlsx"
new_wb.save(s)
print(f"\nФайл успешно сохранен в: {s}")

# Закрытие оригинального файла
wb.close()
